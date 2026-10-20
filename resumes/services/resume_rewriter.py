import io
import os

from docx import Document
from google import genai


def rewrite_resume(resume_text, job_description):
    """
    Resume-ஐ job description-yetha maari maaritu problem ku yeth amaari none return aagum
    
    """

    api_key = os.environ.get("GEMINI_API_KEY")
    if not api_key:
        return None

    prompt = f"""You are an expert resume writer.

Rewrite the resume below so it fits the job description better.

STRICT RULES:
- Use ONLY facts that already exist in the resume.
- Do NOT invent skills, projects, companies, degrees, dates or numbers.
- Do NOT add any skill the candidate has not mentioned.
- Keep name and contact details exactly as given.
- Improve wording: start bullets with strong action verbs, be concise.
- Reorder sections so the most relevant ones for the job come first.
- Output plain text only. Section titles in UPPERCASE on their own line.
  Bullets must start with "- ". No markdown, no asterisks.

JOB DESCRIPTION:
{job_description[:3000]}

RESUME:
{resume_text[:8000]}
"""

    try:
        client = genai.Client(api_key=api_key)
        response = client.models.generate_content(
            model="gemini-flash-latest",
            contents=prompt,
        )
        return response.text.strip()

    except Exception as error:
        print("REWRITE ERROR:", repr(error))
        return None


def build_docx(text):
    """Plain text-ஐ DOCX bytes ஆக மாற்றும்."""

    document = Document()

    for line in text.splitlines():
        line = line.strip()

        if not line:
            continue

        if line.startswith("- "):
            document.add_paragraph(line[2:], style="List Bullet")
        elif line.isupper() and len(line) < 60:
            document.add_heading(line.title(), level=2)
        else:
            document.add_paragraph(line)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()